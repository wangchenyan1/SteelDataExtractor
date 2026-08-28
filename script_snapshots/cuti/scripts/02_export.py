"""Cu-Ti 数据库导出脚本。

从 outputs/<paper_id>/paper.json 直接导出 final_database.csv / xlsx。

设计原则：
- 一行 = 一个 sample × 一个 condition 组合
- 图片直接由 LLM 输出的 figure → sample/condition 映射决定，不再用规则猜测
- 只有 is_microstructure_image=true 且 figure_type ∈ {OM, SEM, TEM} 才进入对应列
- 模板缺列时自动补：硬度/导电率/弯曲、product_form、alloy_family、
  solution_treatment/aging/cold_rolling、overview、比例尺
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_DIR / "outputs"
DATABASE_DIR = PROJECT_DIR / "database_tables"
TEMPLATE_PATH = DATABASE_DIR / "tet.xlsx"


METHODS = ("OM", "SEM", "TEM")
ELEMENTS = (
    "H", "B", "C", "N", "O", "F", "Na", "Mg", "Al", "Si", "P", "S", "Cl",
    "Ca", "Ti", "V", "Cr", "Mn", "Fe", "Co", "Ni", "Cu", "Zn", "As", "Y",
    "Zr", "Nb", "Mo", "Sn", "Sb", "La", "Ce", "Ta", "W", "Pb", "Bi",
)
METHOD_COLUMNS = {
    "OM": ("OM照片编号", "OM照片链接", "OM照片", "OM照片描述"),
    "SEM": ("SEM照片编号", "SEM照片链接", "SEM照片", "SEM照片描述"),
    "TEM": ("TEM照片编号", "TEM照片链接", "TEM照片", "TEM照片描述"),
}
METHOD_SCALE_COLUMNS = {
    "OM": "OM比例尺",
    "SEM": "SEM比例尺",
    "TEM": "TEM比例尺",
}

# (json_key, name_col, value_col, unit_col, 默认 name)
MECH_FIELDS = [
    ("yield_strength",   "Yield_name",      "Yield_value",      "Yield_unit",      "Yield strength"),
    ("tensile_strength", "Tensile_name",    "Tensile_value",    "Tensile_unit",    "Tensile strength"),
    ("elongation",       "Elongation_name", "Elongation_value", "Elongation_unit", "Elongation"),
    ("hardness",         "Hardness_name",   "Hardness_value",   "Hardness_unit",   "Hardness"),
]
ELEC_FIELDS = [
    ("electrical_conductivity", "Conductivity_name", "Conductivity_value", "Conductivity_unit", "Electrical conductivity"),
]


def paper_id_to_doi(paper_id: str) -> str:
    """从 paper_id 反推 DOI。

    paper_id 是 DOI 把 / 换成 _ 的形式，例如：
      10.1007_s10853-007-2000-4 → 10.1007/s10853-007-2000-4
      10.1016_j.actamat.2020.116373 → 10.1016/j.actamat.2020.116373
    """
    s = str(paper_id or "").strip()
    # 只替换第一个 _ 为 /（DOI 前缀 10.XXXX 之后的第一个分隔符）
    m = re.match(r"^(10\.\d{4,9})[_-](.+)$", s)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    return s


def read_template_headers(template: Path) -> list[str]:
    wb = load_workbook(template, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return [h for h in headers if h]


def _ensure_after(headers: list[str], col: str, after: str | None) -> None:
    if col in headers:
        return
    if after and after in headers:
        headers.insert(headers.index(after) + 1, col)
    else:
        headers.append(col)


def ensure_cuti_headers(headers: list[str]) -> list[str]:
    """模板缺列时自动补齐 CuTi 新增字段。"""
    result = list(headers)

    # 力学性能：接在 Elongation_unit 后
    anchor = "Elongation_unit" if "Elongation_unit" in result else None
    for _, name_col, value_col, unit_col, _ in MECH_FIELDS:
        for col in (name_col, value_col, unit_col):
            if col not in result:
                _ensure_after(result, col, anchor)
            anchor = col

    # 导电率
    for _, name_col, value_col, unit_col, _ in ELEC_FIELDS:
        for col in (name_col, value_col, unit_col):
            if col not in result:
                _ensure_after(result, col, anchor)
            anchor = col

    # 弯曲（文本）
    _ensure_after(result, "Bending_performance", anchor)

    # product_form / alloy_family：Material 后
    _ensure_after(result, "product_form", "Material" if "Material" in result else None)
    _ensure_after(result, "alloy_family", "product_form" if "product_form" in result else
                  ("Material" if "Material" in result else None))

    # composition_unit：接在最后一个成分元素列后
    if "composition_unit" not in result:
        last_elem = None
        for elem in reversed(ELEMENTS):
            if elem in result:
                last_elem = elem
                break
        _ensure_after(result, "composition_unit", last_elem)

    # condition 级工艺定格式字段：接在 composition_unit（或 product_form）后
    proc_anchor = "composition_unit" if "composition_unit" in result else (
        "alloy_family" if "alloy_family" in result else "product_form"
    )
    for col in ("solution_treatment", "aging", "cold_rolling"):
        _ensure_after(result, col, proc_anchor if proc_anchor in result else None)
        proc_anchor = col

    # overview
    if "sample_processing_overview" not in result:
        _ensure_after(result, "sample_processing_overview", "Text" if "Text" in result else None)
    if "sample_microstructure_overview" not in result:
        if "组织描述" in result:
            _ensure_after(result, "sample_microstructure_overview", "组织描述")
        else:
            _ensure_after(result, "sample_microstructure_overview", "sample_processing_overview")

    # 比例尺：插在对应照片列后
    for method, cols in METHOD_COLUMNS.items():
        scale_col = METHOD_SCALE_COLUMNS[method]
        if scale_col in result:
            continue
        embed_col = cols[2]
        if embed_col in result:
            result.insert(result.index(embed_col) + 1, scale_col)
        else:
            result.append(scale_col)

    return result


def _set_metric(row: dict, name_col: str, value_col: str, unit_col: str,
                payload: dict | None, default_name: str) -> None:
    if not payload or not isinstance(payload, dict):
        return
    val = payload.get("value", "")
    unit = payload.get("unit", "")
    val_str = "" if val in (None, "") else str(val)
    unit_str = "" if unit in (None, "") else str(unit)
    if val_str or unit_str:
        row[name_col] = default_name
        row[value_col] = val_str
        row[unit_col] = unit_str


def _scale_bar_text(fig: dict) -> str:
    scale = fig.get("scale_bar_info")
    if not isinstance(scale, dict):
        return ""
    value = scale.get("value")
    unit = scale.get("unit")
    value_str = "" if value in (None, "") else str(value)
    unit_str = "" if unit in (None, "") else str(unit)
    if not value_str and not unit_str:
        return ""
    return f"{value_str} {unit_str}".strip()


def materialize_image(image_index: int, paper_id: str, target_dir: Path) -> Path | None:
    """从 outputs/<paper_id>/images_raw 复制到 outputs/<paper_id>/images/<sample>/<method> 目录。

    返回写入路径（用于数据库链接列），不存在返回 None。
    """
    raw_dir = OUTPUT_DIR / paper_id / "images_raw"
    matches = list(raw_dir.glob(f"placeholder_{image_index:03d}.*"))
    if not matches:
        return None
    src = matches[0]
    target_dir.mkdir(parents=True, exist_ok=True)
    dst = target_dir / src.name
    if not dst.exists():
        shutil.copy2(src, dst)
    return dst


def build_rows_for_paper(paper_json_path: Path) -> tuple[list[dict[str, str]], dict]:
    """从 paper.json 生成数据库行列表 + 诊断信息。

    返回 (rows, diagnostics)：
      rows: 字典列表，每个字典是一行
      diagnostics: {paper_id, samples, conditions, figures, microstructure_figures, ...}
    """
    data = json.loads(paper_json_path.read_text(encoding="utf-8"))
    paper_id = data.get("_paper_id") or paper_json_path.parent.name
    metadata = data.get("paper_metadata", {})
    samples = data.get("samples", [])
    conditions = data.get("conditions", [])
    figures = data.get("figures", [])

    # DOI 直接从 paper_id 反推（paper_id 是 DOI 把 / 换成 _ 的形式）
    doi = paper_id_to_doi(paper_id)

    # 索引：sample_id → sample dict
    sample_by_id = {s["sample_id"]: s for s in samples}
    # 索引：sample_id → condition list
    conditions_by_sample: dict[str, list[dict]] = defaultdict(list)
    for c in conditions:
        sid = c.get("sample_id")
        if sid:
            conditions_by_sample[sid].append(c)

    # 索引：figure 按 (sample_id, condition_id, figure_type) 分桶
    # 注意：严格过滤 post-test 图（is_post_test_image=true 或 image_purpose 是 post_test_*）
    figures_by_key: dict[tuple[str, str | None, str], list[dict]] = defaultdict(list)
    for fig in figures:
        if not fig.get("is_microstructure_image"):
            continue
        if fig.get("is_post_test_image"):
            continue
        purpose = (fig.get("image_purpose") or "").lower()
        if purpose in {"post_test_fractography", "fracture_surface", "mechanical_curve", "diffraction", "schematic"}:
            continue
        ftype = (fig.get("figure_type") or "").upper()
        if ftype not in METHODS:
            continue
        # caption 兜底过滤（即使 LLM 没标 is_post_test）
        caption_lower = (fig.get("caption") or "").lower()
        if any(kw in caption_lower for kw in ("fractured sample", "post-test", "post test",
                                                "after tensile", "interrupted tensile",
                                                "deformed to fracture")):
            continue
        key = (fig.get("sample_id") or "", fig.get("condition_id"), ftype)
        figures_by_key[key].append(fig)

    rows: list[dict[str, str]] = []

    for sample in samples:
        sid = sample["sample_id"]
        sname = sample.get("sample_name", sid)
        composition = sample.get("composition", {}) or {}
        base_proc = sample.get("base_processing_description") or sample.get("processing_description") or ""
        product_form = sample.get("product_form") or ""
        alloy_family = sample.get("alloy_family") or ""
        sample_proc_overview = sample.get("sample_processing_overview") or ""
        sample_micro_overview = sample.get("sample_microstructure_overview") or ""
        sample_conditions = conditions_by_sample.get(sid) or []

        # 没有 condition 的样品：建一行最小条目（仅成分）
        if not sample_conditions:
            sample_conditions = [None]

        for cond in sample_conditions:
            # 过滤 post_test_condition：不导出（避免断后样品入表）
            if cond is not None:
                ctype = (cond.get("condition_type") or "").lower()
                if ctype == "post_test_condition":
                    continue

            cond_id_for_uid = (cond.get("condition_id") if cond else "") or "C0"
            row: dict[str, str] = {
                "UID": f"{paper_id}-{sid}-{cond_id_for_uid}",
                "DOIs": doi,
            }

            # Material 列：使用 condition_name（工艺态代号），不是成分式
            cond_name = (cond.get("condition_name") if cond else "") or ""
            row["Material"] = cond_name or sname  # 无 condition 时退回 sample 名
            row["product_form"] = str(product_form) if product_form not in ("", None) else ""
            row["alloy_family"] = str(alloy_family) if alloy_family not in ("", None) else ""

            # Table_topic：sample 名 + 工艺态简称
            row["Table_topic"] = f"{sname} | {cond_name}" if cond_name else sname

            # 成分列
            for elem in ELEMENTS:
                val = composition.get(elem, "")
                row[elem] = str(val) if val not in ("", None) else ""
            unit = sample.get("composition_unit", "")
            row["composition_unit"] = "" if unit in ("", None) else str(unit)

            # condition 级工艺定格式字段
            for col in ("solution_treatment", "aging", "cold_rolling"):
                val = cond.get(col) if cond else None
                row[col] = str(val) if val not in (None, "") else ""

            # 性能列（按 condition）
            mech = (cond.get("mechanical_properties", {}) if cond else {}) or {}
            for json_key, name_col, value_col, unit_col, default_name in MECH_FIELDS:
                _set_metric(row, name_col, value_col, unit_col, mech.get(json_key), default_name)

            elec = (cond.get("electrical_properties", {}) if cond else {}) or {}
            for json_key, name_col, value_col, unit_col, default_name in ELEC_FIELDS:
                _set_metric(row, name_col, value_col, unit_col, elec.get(json_key), default_name)

            bend = cond.get("bending_performance") if cond else None
            if bend not in (None, ""):
                row["Bending_performance"] = str(bend)

            # 组织描述（基础值；OM/SEM/TEM 图注会在图片循环后追加）
            micro_desc = cond.get("microstructure_description", "") if cond else ""

            # Text 字段：sample 名 + condition 名 + base 工艺 + condition 独有工艺
            cond_proc = cond.get("condition_processing_description", "") if cond else ""
            text_parts = [sname]
            if cond_name:
                text_parts.append(cond_name)
            if alloy_family:
                text_parts.append(alloy_family)
            for col in ("solution_treatment", "aging", "cold_rolling"):
                val = row.get(col) or ""
                if val:
                    text_parts.append(f"{col}={val}")
            if sample_proc_overview:
                text_parts.append(sample_proc_overview)
            elif base_proc:
                text_parts.append(base_proc)
            if cond_proc:
                text_parts.append(cond_proc)
            row["Text"] = " | ".join(text_parts)
            row["sample_processing_overview"] = sample_proc_overview
            row["sample_microstructure_overview"] = sample_micro_overview

            method_captions: dict[str, list[str]] = {}

            # 图片列：按 sample_id + condition_id + method 查找
            cond_id = cond.get("condition_id") if cond else None
            for method in METHODS:
                # 精确匹配 condition
                figs = figures_by_key.get((sid, cond_id, method), [])
                # 若无,再用 condition_id=None 的 sample-level 图
                if not figs:
                    figs = figures_by_key.get((sid, None, method), [])
                if not figs:
                    continue

                num_col, link_col, embed_col, desc_col = METHOD_COLUMNS[method]
                scale_col = METHOD_SCALE_COLUMNS[method]
                links: list[str] = []
                numbers: list[str] = []
                descs: list[str] = []
                scales: list[str] = []
                for fig in figs:
                    idx = fig.get("placeholder_index")
                    if not idx:
                        continue
                    sample_safe = sname.replace("/", "_").replace(" ", "_")
                    target = OUTPUT_DIR / paper_id / "images" / sample_safe / method.lower()
                    dst = materialize_image(idx, paper_id, target)
                    if not dst:
                        continue
                    rel = dst.relative_to(PROJECT_DIR).as_posix()
                    links.append(rel)
                    numbers.append(fig.get("figure_id", ""))
                    descs.append(fig.get("caption", ""))
                    scales.append(_scale_bar_text(fig))

                if links:
                    row[num_col] = "; ".join(filter(None, numbers))
                    row[link_col] = "; ".join(links)
                    row[embed_col] = links[0]  # 第一张用于 XLSX 嵌入
                    row[scale_col] = "; ".join(s for s in scales if s)
                    row[desc_col] = "; ".join(filter(None, descs))
                    method_captions[method] = [d for d in descs if d]

            # 组织描述：基础描述 + 各方法图注
            micro_parts: list[str] = []
            if micro_desc:
                micro_parts.append(micro_desc)
            for method in METHODS:
                caps = method_captions.get(method)
                if caps:
                    micro_parts.append(f"{method}: " + "; ".join(caps))
            row["组织描述"] = " | ".join(micro_parts)

            rows.append(row)

    diagnostics = {
        "paper_id": paper_id,
        "samples_count": len(samples),
        "conditions_count": len(conditions),
        "figures_total": len(figures),
        "microstructure_figures": sum(1 for f in figures if f.get("is_microstructure_image") and not f.get("is_post_test_image")),
        "post_test_figures": sum(1 for f in figures if f.get("is_post_test_image")),
        "rows_count": len(rows),
    }
    return rows, diagnostics


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({h: row.get(h, "") for h in headers})


def set_image_size(image: XlsxImage) -> None:
    max_w, max_h = 200, 140
    try:
        w, h = image.width, image.height
        scale = min(max_w / w, max_h / h, 1.0)
        image.width = int(w * scale)
        image.height = int(h * scale)
    except Exception:
        pass


def write_xlsx(path: Path, headers: list[str], rows: list[dict[str, str]], root: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "final_database"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    image_cols = {"OM照片", "SEM照片", "TEM照片"}
    image_col_idx = {h: headers.index(h) + 1 for h in image_cols if h in headers}
    inserted = 0
    for row_index, row in enumerate(rows, start=2):
        ws.append([row.get(h, "") for h in headers])
        ws.row_dimensions[row_index].height = 115
        for h, col_idx in image_col_idx.items():
            rel = row.get(h, "").strip()
            if not rel:
                continue
            image_path = root / rel
            if not image_path.exists():
                continue
            try:
                img = XlsxImage(str(image_path))
                set_image_size(img)
                anchor = f"{get_column_letter(col_idx)}{row_index}"
                ws.add_image(img, anchor)
                inserted += 1
            except Exception as exc:
                print(f"[WARN] 插入图片失败 {rel}: {exc}")

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if header in image_cols:
            ws.column_dimensions[letter].width = 28
        elif (
            "描述" in header
            or header == "Text"
            or header in (
                "sample_processing_overview",
                "sample_microstructure_overview",
                "Bending_performance",
            )
        ):
            ws.column_dimensions[letter].width = 45
        elif header in ("solution_treatment", "aging", "cold_rolling", "alloy_family"):
            ws.column_dimensions[letter].width = 28
        elif "链接" in header:
            ws.column_dimensions[letter].width = 42
        elif header == "UID":
            ws.column_dimensions[letter].width = 42
        else:
            ws.column_dimensions[letter].width = min(max(len(header) + 4, 12), 24)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return inserted


def write_summary(path: Path, diagnostics: list[dict], rows: list[dict[str, str]], inserted: int) -> None:
    lines = [
        "# V4 Export Summary",
        "",
        f"- total papers: {len(diagnostics)}",
        f"- total rows: {len(rows)}",
        f"- xlsx inserted images: {inserted}",
        "",
        "## Per-Paper Diagnostics",
        "",
        "| paper_id | samples | conditions | figures | microstructure | rows |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for d in diagnostics:
        lines.append(
            f"| {d['paper_id']} | {d['samples_count']} | {d['conditions_count']} | "
            f"{d['figures_total']} | {d['microstructure_figures']} | {d['rows_count']} |"
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="V4 数据库导出")
    parser.add_argument("--paper-id", action="append", help="仅导出指定论文（可多次）")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--output-prefix", default="final_database",
                        help="输出文件名前缀（默认 final_database）")
    args = parser.parse_args()

    if not TEMPLATE_PATH.exists():
        print(f"[ERROR] 模板不存在: {TEMPLATE_PATH}", file=sys.stderr)
        return 1
    headers = read_template_headers(TEMPLATE_PATH)
    if "UID" not in headers:
        headers = ["UID", *headers]
    headers = ensure_cuti_headers(headers)
    print(f"[INFO] 模板列数: {len(headers)}")

    if args.paper_id:
        paper_ids = args.paper_id
    else:
        paper_ids = sorted(
            d.name for d in OUTPUT_DIR.iterdir()
            if d.is_dir() and (d / "paper.json").exists()
        )
    print(f"[INFO] 待导出论文数: {len(paper_ids)}")

    all_rows: list[dict[str, str]] = []
    diagnostics: list[dict] = []
    for pid in paper_ids:
        pj = OUTPUT_DIR / pid / "paper.json"
        if not pj.exists():
            print(f"[SKIP] {pid}: paper.json 缺失")
            continue
        try:
            rows, diag = build_rows_for_paper(pj)
            all_rows.extend(rows)
            diagnostics.append(diag)
            print(f"[OK] {pid}: {len(rows)} 行 / {diag['samples_count']} 样品 / "
                  f"{diag['microstructure_figures']} 组织图")
        except Exception as exc:
            print(f"[FAIL] {pid}: {exc}")
            import traceback
            traceback.print_exc()

    csv_path = DATABASE_DIR / f"{args.output_prefix}.csv"
    xlsx_path = DATABASE_DIR / f"{args.output_prefix}.xlsx"
    summary_path = DATABASE_DIR / f"{args.output_prefix}_summary.md"

    if csv_path.exists() and not args.overwrite:
        print(f"[ERROR] 输出文件已存在，使用 --overwrite 覆盖: {csv_path}")
        return 1

    write_csv(csv_path, headers, all_rows)
    inserted = write_xlsx(xlsx_path, headers, all_rows, PROJECT_DIR)
    write_summary(summary_path, diagnostics, all_rows, inserted)

    print(f"\n[OK] CSV: {csv_path}")
    print(f"[OK] XLSX: {xlsx_path} ({inserted} 张图片嵌入)")
    print(f"[OK] Summary: {summary_path}")
    print(f"[INFO] 总计: {len(diagnostics)} 篇 / {len(all_rows)} 行")
    return 0


if __name__ == "__main__":
    sys.exit(main())
