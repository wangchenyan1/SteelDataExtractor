"""无磁钢专用：数据库导出脚本。

从 无磁钢/outputs/<paper_id>/paper.json 导出 final_database.csv / xlsx。

设计原则：
- 一行 = 一个 sample × 一个 condition 组合
- 图片由 LLM 输出的 figure → sample/condition 映射决定
- 只有 is_microstructure_image=true 且 figure_type ∈ {OM, SEM, TEM} 才进入对应列
- 输出列与 无磁钢/database_tables/tet.xlsx 模板一致
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parents[2]  # paper_processing_v4/
SUBPROJECT_DIR = Path(__file__).resolve().parents[1]  # 无磁钢/
OUTPUT_DIR = SUBPROJECT_DIR / "outputs"
DATABASE_DIR = SUBPROJECT_DIR / "database_tables"
TEMPLATE_PATH = DATABASE_DIR / "tet.xlsx"


METHODS = ("OM", "SEM", "TEM")
ELEMENTS = ("Fe", "C", "Si", "Mn", "P", "S", "Cr", "Ni", "Mo", "N", "Nb", "V", "Cu", "Al")

METHOD_COLUMNS = {
    "OM":  ("OM照片编号", "OM照片链接", "OM照片", "OM比例尺", "OM照片描述"),
    "SEM": ("SEM照片编号", "SEM照片链接", "SEM照片", "SEM比例尺", "SEM照片描述"),
    "TEM": ("TEM照片编号", "TEM照片链接", "TEM照片", "TEM比例尺", "TEM照片描述"),
}

# (json_key 路径, name_col, value_col, unit_col, 默认 name 字符串)
MECH_FIELDS = [
    ("yield_strength",   "Yield_name",          "Yield_value",          "Yield_unit",          "Yield strength (RT)"),
    ("tensile_strength", "Tensile_name",        "Tensile_value",        "Tensile_unit",        "Tensile strength (RT)"),
    ("elongation",       "Elongation_name",     "Elongation_value",     "Elongation_unit",     "Elongation (RT)"),
    ("elastic_modulus",  "ElasticModulus_name", "ElasticModulus_value", "ElasticModulus_unit", "Elastic modulus"),
    ("hardness",         "Hardness_name",       "Hardness_value",       "Hardness_unit",       "Hardness"),
    ("impact_77K",       "Impact77K_name",      "Impact77K_value",      "Impact77K_unit",      "Impact energy (77K)"),
]
MAG_FIELDS = [
    ("permeability",              "Permeability_name",       "Permeability_value",       "Permeability_unit",       "Permeability"),
    ("remanence",                 "Remanence_name",          "Remanence_value",          "Remanence_unit",          "Remanence"),
    ("saturation_magnetization",  "SatMagnetization_name",   "SatMagnetization_value",   "SatMagnetization_unit",   "Saturation magnetization"),
    ("coercivity",                "Coercivity_name",         "Coercivity_value",         "Coercivity_unit",         "Coercivity"),
    ("curie_temperature",         "CurieTemp_name",          "CurieTemp_value",          "CurieTemp_unit",          "Curie temperature"),
    ("neel_temperature",          "NeelTemp_name",           "NeelTemp_value",           "NeelTemp_unit",           "Neel temperature"),
]
# 固定格式字符串型磁性能字段：(json_key, excel_col)
MAG_TEXT_FIELDS = [
    ("austenite_stability", "AusteniteStability"),
    ("martensite_transformation_temperature", "MartensiteTransTemp"),
]
ROLL_FIELDS = [
    ("heating_temperature",        "HeatingTemp_name",       "HeatingTemp_value",       "HeatingTemp_unit",       "Heating temperature"),
    ("start_rolling_temperature",  "StartRollingTemp_name",  "StartRollingTemp_value",  "StartRollingTemp_unit",  "Start rolling temperature"),
    ("finish_rolling_temperature", "FinishRollingTemp_name", "FinishRollingTemp_value", "FinishRollingTemp_unit", "Finish rolling temperature"),
    ("reduction_ratio",            "ReductionRatio_name",    "ReductionRatio_value",    "ReductionRatio_unit",    "Reduction ratio"),
    ("cooling_rate",               "CoolingRate_name",       "CoolingRate_value",       "CoolingRate_unit",       "Cooling rate"),
]


def paper_id_to_doi(paper_id: str) -> str:
    s = str(paper_id or "").strip()
    m = re.match(r"^(10\.\d{4,9})[_-](.+)$", s)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    return s


def read_template_headers(template: Path) -> list[str]:
    wb = load_workbook(template, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return [h for h in headers if h]


def ensure_magnetic_headers(headers: list[str]) -> list[str]:
    """模板缺少新增磁性能列时，按顺序插在已有磁性能列之后（或追加到末尾）。"""
    result = list(headers)
    anchor = None
    for col in ("Remanence_unit", "Remanence_value", "Permeability_unit"):
        if col in result:
            anchor = col
            break

    def _ensure(col: str) -> None:
        nonlocal anchor
        if col in result:
            anchor = col
            return
        if anchor and anchor in result:
            result.insert(result.index(anchor) + 1, col)
        else:
            result.append(col)
        anchor = col

    for _, name_col, value_col, unit_col, _ in MAG_FIELDS:
        for col in (name_col, value_col, unit_col):
            _ensure(col)
    for _, excel_col in MAG_TEXT_FIELDS:
        _ensure(excel_col)
    return result


EXTRA_EXPORT_COLUMNS = [
    "ProductForm",
    "AlloyFamily",
    "SampleProcessingOverview",
    "SampleMicrostructureOverview",
    "OM比例尺",
    "SEM比例尺",
    "TEM比例尺",
]

def _insert_before(result: list[str], col: str, anchor: str) -> None:
    if col in result:
        return
    if anchor in result:
        result.insert(result.index(anchor), col)
    else:
        result.append(col)


def _insert_after(result: list[str], col: str, anchor: str) -> None:
    if col in result:
        return
    if anchor in result:
        result.insert(result.index(anchor) + 1, col)
    else:
        result.append(col)


def ensure_extra_headers(headers: list[str]) -> list[str]:
    """补齐新增抽取字段，并尽量放到模板中语义相邻的位置。"""
    result = list(headers)
    for col in ("ProductForm", "AlloyFamily"):
        if col not in result:
            result.append(col)
    # 两个 overview 是组织相关文本，放在「组织描述」前面。
    _insert_before(result, "SampleProcessingOverview", "组织描述")
    _insert_before(result, "SampleMicrostructureOverview", "组织描述")
    # 比例尺跟随对应图片，放在「照片」和「照片描述」之间。
    for method in METHODS:
        _insert_after(result, f"{method}比例尺", f"{method}照片")
    return result


def materialize_image(image_index: int, paper_id: str, target_dir: Path) -> Path | None:
    raw_dir = OUTPUT_DIR / paper_id / "images_raw"
    matches = list(raw_dir.glob(f"placeholder_{image_index:03d}.*"))
    if not matches:
        return None
    src = matches[0]
    target_dir.mkdir(parents=True, exist_ok=True)
    dst = target_dir / src.name
    if not dst.exists():
        shutil.copy2(src, dst)
    return dst


def _set_metric(row: dict, name_col: str, value_col: str, unit_col: str,
                payload: dict | None, default_name: str) -> None:
    """把 {"value": "650", "unit": "MPa"} 形态的字典写到三列。"""
    if not payload or not isinstance(payload, dict):
        return
    val = payload.get("value", "")
    unit = payload.get("unit", "")
    val_str = "" if val in (None, "") else str(val)
    unit_str = "" if unit in (None, "") else str(unit)
    if val_str or unit_str:
        row[name_col] = default_name
        row[value_col] = val_str
        row[unit_col] = unit_str


def format_scale_bar(scale: object) -> str:
    if isinstance(scale, dict):
        value = scale.get("value")
        unit = scale.get("unit")
        if value in (None, "") and unit in (None, ""):
            return ""
        return f"{value or ''} {unit or ''}".strip()
    if scale in (None, ""):
        return ""
    return str(scale)


def build_rows_for_paper(paper_json_path: Path) -> tuple[list[dict[str, str]], dict]:
    data = json.loads(paper_json_path.read_text(encoding="utf-8"))
    paper_id = data.get("_paper_id") or paper_json_path.parent.name
    samples = data.get("samples", [])
    conditions = data.get("conditions", [])
    figures = data.get("figures", [])
    doi = paper_id_to_doi(paper_id)

    conditions_by_sample: dict[str, list[dict]] = defaultdict(list)
    for c in conditions:
        sid = c.get("sample_id")
        if sid:
            conditions_by_sample[sid].append(c)

    figures_by_key: dict[tuple[str, str | None, str], list[dict]] = defaultdict(list)
    for fig in figures:
        if not fig.get("is_microstructure_image"):
            continue
        if fig.get("is_post_test_image"):
            continue
        purpose = (fig.get("image_purpose") or "").lower()
        if purpose in {"post_test_fractography", "fracture_surface", "mechanical_curve",
                       "magnetic_curve", "diffraction", "schematic"}:
            continue
        ftype = (fig.get("figure_type") or "").upper()
        if ftype not in METHODS:
            continue
        caption_lower = (fig.get("caption") or "").lower()
        if any(kw in caption_lower for kw in ("fractured sample", "post-test", "post test",
                                              "after tensile", "interrupted tensile",
                                              "deformed to fracture")):
            continue
        key = (fig.get("sample_id") or "", fig.get("condition_id"), ftype)
        figures_by_key[key].append(fig)

    rows: list[dict[str, str]] = []

    for sample in samples:
        sid = sample["sample_id"]
        sname = sample.get("sample_name", sid)
        composition = sample.get("composition", {}) or {}
        base_proc = sample.get("base_processing_description") or sample.get("processing_description") or ""
        sample_conditions = conditions_by_sample.get(sid) or [None]

        for cond in sample_conditions:
            if cond is not None:
                ctype = (cond.get("condition_type") or "").lower()
                if ctype == "post_test_condition":
                    continue

            cond_id_for_uid = (cond.get("condition_id") if cond else "") or "C0"
            row: dict[str, str] = {
                "UID": f"{paper_id}-{sid}-{cond_id_for_uid}",
                "DOIs": doi,
                "ProductForm": str(sample.get("product_form") or ""),
                "AlloyFamily": str(sample.get("alloy_family") or ""),
                "SampleProcessingOverview": str(sample.get("sample_processing_overview") or ""),
                "SampleMicrostructureOverview": str(sample.get("sample_microstructure_overview") or ""),
            }

            cond_name = (cond.get("condition_name") if cond else "") or ""
            row["Material"] = cond_name or sname
            row["Table_topic"] = f"{sname} | {cond_name}" if cond_name else sname

            for elem in ELEMENTS:
                val = composition.get(elem, "")
                row[elem] = str(val) if val not in ("", None) else ""
            unit = sample.get("composition_unit", "")
            row["composition_unit"] = str(unit) if unit not in ("", None) else ""

            mech = (cond.get("mechanical_properties", {}) if cond else {}) or {}
            for json_key, name_col, value_col, unit_col, default_name in MECH_FIELDS:
                _set_metric(row, name_col, value_col, unit_col, mech.get(json_key), default_name)

            mag = (cond.get("magnetic_properties", {}) if cond else {}) or {}
            for json_key, name_col, value_col, unit_col, default_name in MAG_FIELDS:
                _set_metric(row, name_col, value_col, unit_col, mag.get(json_key), default_name)
            for json_key, excel_col in MAG_TEXT_FIELDS:
                text_val = mag.get(json_key)
                if text_val not in (None, ""):
                    row[excel_col] = str(text_val)

            roll = (cond.get("rolling_processing", {}) if cond else {}) or {}
            for json_key, name_col, value_col, unit_col, default_name in ROLL_FIELDS:
                _set_metric(row, name_col, value_col, unit_col, roll.get(json_key), default_name)

            micro = (cond.get("microstructure", {}) if cond else {}) or {}
            _set_metric(row, "GrainSize_name", "GrainSize_value", "GrainSize_unit",
                        micro.get("grain_size"), "Grain size")
            micro_desc = (micro.get("description") or "") if isinstance(micro, dict) else ""

            cond_proc = cond.get("condition_processing_description", "") if cond else ""
            text_parts = [sname]
            if cond_name:
                text_parts.append(cond_name)
            if base_proc:
                text_parts.append(base_proc)
            if cond_proc:
                text_parts.append(cond_proc)
            row["Text"] = " | ".join(text_parts)

            method_captions: dict[str, list[str]] = {}

            cond_id = cond.get("condition_id") if cond else None
            for method in METHODS:
                figs = figures_by_key.get((sid, cond_id, method), [])
                if not figs:
                    figs = figures_by_key.get((sid, None, method), [])
                if not figs:
                    continue

                num_col, link_col, embed_col, scale_col, desc_col = METHOD_COLUMNS[method]
                links: list[str] = []
                numbers: list[str] = []
                descs: list[str] = []
                scales: list[str] = []
                for fig in figs:
                    idx = fig.get("placeholder_index")
                    if not idx:
                        continue
                    sample_safe = sname.replace("/", "_").replace(" ", "_")
                    target = OUTPUT_DIR / paper_id / "images" / sample_safe / method.lower()
                    dst = materialize_image(idx, paper_id, target)
                    if not dst:
                        continue
                    rel = dst.relative_to(PROJECT_DIR).as_posix()
                    links.append(rel)
                    numbers.append(fig.get("figure_id", ""))
                    descs.append(fig.get("caption", ""))
                    scales.append(format_scale_bar(fig.get("scale_bar_info")))

                if links:
                    row[num_col] = "; ".join(filter(None, numbers))
                    row[link_col] = "; ".join(links)
                    row[embed_col] = links[0]
                    row[scale_col] = "; ".join(filter(None, scales))
                    row[desc_col] = "; ".join(filter(None, descs))
                    method_captions[method] = [d for d in descs if d]

            micro_parts: list[str] = []
            if micro_desc:
                micro_parts.append(micro_desc)
            for method in METHODS:
                caps = method_captions.get(method)
                if caps:
                    micro_parts.append(f"{method}: " + "; ".join(caps))
            row["组织描述"] = " | ".join(micro_parts)

            rows.append(row)

    diagnostics = {
        "paper_id": paper_id,
        "samples_count": len(samples),
        "conditions_count": len(conditions),
        "figures_total": len(figures),
        "microstructure_figures": sum(1 for f in figures if f.get("is_microstructure_image") and not f.get("is_post_test_image")),
        "post_test_figures": sum(1 for f in figures if f.get("is_post_test_image")),
        "rows_count": len(rows),
    }
    return rows, diagnostics


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({h: row.get(h, "") for h in headers})


def set_image_size(image: XlsxImage) -> None:
    max_w, max_h = 200, 140
    try:
        w, h = image.width, image.height
        scale = min(max_w / w, max_h / h, 1.0)
        image.width = int(w * scale)
        image.height = int(h * scale)
    except Exception:
        pass


def write_xlsx(path: Path, headers: list[str], rows: list[dict[str, str]], root: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "final_database"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    image_cols = {"OM照片", "SEM照片", "TEM照片"}
    image_col_idx = {h: headers.index(h) + 1 for h in image_cols if h in headers}
    inserted = 0
    for row_index, row in enumerate(rows, start=2):
        ws.append([row.get(h, "") for h in headers])
        ws.row_dimensions[row_index].height = 115
        for h, col_idx in image_col_idx.items():
            rel = row.get(h, "").strip()
            if not rel:
                continue
            image_path = root / rel
            if not image_path.exists():
                continue
            try:
                img = XlsxImage(str(image_path))
                set_image_size(img)
                anchor = f"{get_column_letter(col_idx)}{row_index}"
                ws.add_image(img, anchor)
                inserted += 1
            except Exception as exc:
                print(f"[WARN] 插入图片失败 {rel}: {exc}")

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if header in image_cols:
            ws.column_dimensions[letter].width = 28
        elif "描述" in header or header == "Text":
            ws.column_dimensions[letter].width = 45
        elif "链接" in header:
            ws.column_dimensions[letter].width = 42
        elif header == "UID":
            ws.column_dimensions[letter].width = 42
        else:
            ws.column_dimensions[letter].width = min(max(len(header) + 4, 12), 24)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return inserted


def write_summary(path: Path, diagnostics: list[dict], rows: list[dict[str, str]], inserted: int) -> None:
    lines = [
        "# 无磁钢 Export Summary",
        "",
        f"- total papers: {len(diagnostics)}",
        f"- total rows: {len(rows)}",
        f"- xlsx inserted images: {inserted}",
        "",
        "## Per-Paper Diagnostics",
        "",
        "| paper_id | samples | conditions | figures | microstructure | rows |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for d in diagnostics:
        lines.append(
            f"| {d['paper_id']} | {d['samples_count']} | {d['conditions_count']} | "
            f"{d['figures_total']} | {d['microstructure_figures']} | {d['rows_count']} |"
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="无磁钢 数据库导出")
    parser.add_argument("--paper-id", action="append", help="仅导出指定论文（可多次）")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--output-prefix", default="final_database",
                        help="输出文件名前缀（默认 final_database）")
    args = parser.parse_args()

    if not TEMPLATE_PATH.exists():
        print(f"[ERROR] 模板不存在: {TEMPLATE_PATH}", file=sys.stderr)
        return 1
    headers = read_template_headers(TEMPLATE_PATH)
    if "UID" not in headers:
        headers = ["UID", *headers]
    if "composition_unit" not in headers:
        # 插在成分元素列之前；若模板无 Fe 列则追加到末尾
        if "Fe" in headers:
            fe_idx = headers.index("Fe")
            headers = [*headers[:fe_idx], "composition_unit", *headers[fe_idx:]]
        else:
            headers = [*headers, "composition_unit"]
    headers = ensure_magnetic_headers(headers)
    headers = ensure_extra_headers(headers)
    print(f"[INFO] 模板列数: {len(headers)}")

    if args.paper_id:
        paper_ids = args.paper_id
    else:
        paper_ids = sorted(
            d.name for d in OUTPUT_DIR.iterdir()
            if d.is_dir() and (d / "paper.json").exists()
        )
    print(f"[INFO] 待导出论文数: {len(paper_ids)}")

    all_rows: list[dict[str, str]] = []
    diagnostics: list[dict] = []
    for pid in paper_ids:
        pj = OUTPUT_DIR / pid / "paper.json"
        if not pj.exists():
            print(f"[SKIP] {pid}: paper.json 缺失")
            continue
        try:
            rows, diag = build_rows_for_paper(pj)
            all_rows.extend(rows)
            diagnostics.append(diag)
            print(f"[OK] {pid}: {len(rows)} 行 / {diag['samples_count']} 样品 / "
                  f"{diag['microstructure_figures']} 组织图")
        except Exception as exc:
            print(f"[FAIL] {pid}: {exc}")
            import traceback
            traceback.print_exc()

    csv_path = DATABASE_DIR / f"{args.output_prefix}.csv"
    xlsx_path = DATABASE_DIR / f"{args.output_prefix}.xlsx"
    summary_path = DATABASE_DIR / f"{args.output_prefix}_summary.md"

    if csv_path.exists() and not args.overwrite:
        print(f"[ERROR] 输出文件已存在，使用 --overwrite 覆盖: {csv_path}")
        return 1

    write_csv(csv_path, headers, all_rows)
    inserted = write_xlsx(xlsx_path, headers, all_rows, PROJECT_DIR)
    write_summary(summary_path, diagnostics, all_rows, inserted)

    print(f"\n[OK] CSV: {csv_path}")
    print(f"[OK] XLSX: {xlsx_path} ({inserted} 张图片嵌入)")
    print(f"[OK] Summary: {summary_path}")
    print(f"[INFO] 总计: {len(diagnostics)} 篇 / {len(all_rows)} 行")
    return 0


if __name__ == "__main__":
    sys.exit(main())
